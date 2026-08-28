"""Lead tags: import assignment, Manage Tags, write paths, backup, script groups."""

import json
from pathlib import Path
from unittest.mock import Mock, patch

from django.test import TestCase, override_settings
from django.urls import reverse

from leads.models import CategoryRule, Lead, Tag
from leads.pipeline import get_or_create_uncategorized_group
from leads.services import classify_category_from_name, matching_category_slugs_from_name
from leads.tests import staff_client


def _table_save_data(overrides=None):
    """POST body matching the Manage Tags table (document order)."""
    overrides = overrides or {}
    phrases_by_slug: dict[str, list[str]] = {}
    for category, phrase in CategoryRule.objects.order_by("id").values_list(
        "category", "match_phrase"
    ):
        phrases_by_slug.setdefault(category, []).append(phrase)

    data = {
        "id": [],
        "label": [],
        "slug": [],
        "sort_order": [],
        "match_phrase": [],
    }
    for tag in Tag.objects.order_by("sort_order", "label", "slug"):
        row = overrides.get(tag.pk, {})
        data["id"].append(str(tag.pk))
        data["label"].append(row.get("label", tag.label))
        data["slug"].append(row.get("slug", tag.slug))
        data["sort_order"].append(str(row.get("sort_order", tag.sort_order)))
        if "match_phrase" in row:
            data["match_phrase"].append(row["match_phrase"])
        else:
            existing = phrases_by_slug.get(tag.slug, [])
            data["match_phrase"].append(
                ", ".join(existing) if existing else tag.label
            )
    return data


class NewLeadTagDefaultsTests(TestCase):
    def test_new_lead_starts_with_no_tags(self):
        lead = Lead.objects.create(
            name="No Dual Write Clinic",
            address="1 Tag Test St",
        )
        self.assertEqual(lead.tags.count(), 0)
        self.assertEqual(lead.primary_tag_slug, "unknown")


def _serper_places_response(places):
    resp = Mock()
    resp.raise_for_status = Mock()
    resp.json.return_value = {"places": places}
    return resp


class SerperImportTagTests(TestCase):
    def setUp(self):
        CategoryRule.objects.create(match_phrase="dental", category="dental", priority=10)
        CategoryRule.objects.create(match_phrase="skin", category="aesthetic", priority=20)
        CategoryRule.objects.create(match_phrase="clinic", category="gp", priority=50)

    def test_classify_still_returns_first_priority_match_only(self):
        self.assertEqual(
            classify_category_from_name("Q & M Dental Clinic"),
            "dental",
        )
        self.assertEqual(
            matching_category_slugs_from_name("Q & M Dental Clinic"),
            ["dental", "gp"],
        )
        self.assertEqual(classify_category_from_name("Random Bakery"), "unknown")
        self.assertEqual(matching_category_slugs_from_name("Random Bakery"), [])

    @override_settings(SERPER_API_KEY="test-key", HUNT_MAX_LIMIT=100)
    @patch("leads.services.requests.post")
    def test_new_lead_gets_all_matching_tags_and_single_category(self, mock_post):
        from leads.services import fetch_leads_from_serper

        mock_post.return_value = _serper_places_response(
            [
                {
                    "title": "Q & M Dental Clinic (Segamat)",
                    "address": "1 Tag Import St",
                    "phoneNumber": "+60121111001",
                },
            ]
        )
        result = fetch_leads_from_serper(
            "Segamat",
            "",
            num=20,
            shop_keyword="clinic",
            country="Malaysia",
        )
        self.assertEqual(result.created, 1)
        lead = Lead.objects.get(name="Q & M Dental Clinic (Segamat)")
        self.assertCountEqual(
            list(lead.tags.values_list("slug", flat=True)),
            ["dental", "gp"],
        )
        self.assertEqual(lead.primary_tag_slug, "dental")

    @override_settings(SERPER_API_KEY="test-key", HUNT_MAX_LIMIT=100)
    @patch("leads.services.requests.post")
    def test_single_or_no_rule_match_is_one_tag_matching_category(self, mock_post):
        from leads.services import fetch_leads_from_serper

        mock_post.return_value = _serper_places_response(
            [
                {
                    "title": "Skin Lab Studio",
                    "address": "2 Tag Import St",
                    "phoneNumber": "+60121111002",
                },
                {
                    "title": "Sunrise Hardware",
                    "address": "3 Tag Import St",
                    "phoneNumber": "+60121111003",
                },
            ]
        )
        result = fetch_leads_from_serper(
            "Johor Bahru",
            "",
            num=20,
            shop_keyword="shop",
            country="Malaysia",
        )
        self.assertEqual(result.created, 2)

        single = Lead.objects.get(name="Skin Lab Studio")
        self.assertEqual(single.primary_tag_slug, "aesthetic")
        self.assertEqual(list(single.tags.values_list("slug", flat=True)), ["aesthetic"])

        none = Lead.objects.get(name="Sunrise Hardware")
        self.assertEqual(none.primary_tag_slug, "unknown")
        self.assertEqual(list(none.tags.values_list("slug", flat=True)), ["unknown"])

    @override_settings(SERPER_API_KEY="test-key", HUNT_MAX_LIMIT=100)
    @patch("leads.services.requests.post")
    def test_reimport_leaves_existing_tags_untouched(self, mock_post):
        from leads.services import fetch_leads_from_serper

        existing = Lead.objects.create(
            name="Q & M Dental Clinic (Segamat)",
            address="1 Tag Import St",
        )
        existing.tags.set(Tag.objects.filter(slug="invalid"))
        prior_tag_ids = list(existing.tags.values_list("id", flat=True))

        mock_post.return_value = _serper_places_response(
            [
                {
                    "title": "Q & M Dental Clinic (Segamat)",
                    "address": "1 Tag Import St",
                    "phoneNumber": "+60121111004",
                },
            ]
        )
        result = fetch_leads_from_serper(
            "Segamat",
            "",
            num=20,
            shop_keyword="clinic",
            country="Malaysia",
        )
        self.assertEqual(result.created, 0)
        self.assertEqual(result.skipped_existing, 1)
        existing.refresh_from_db()
        self.assertEqual(list(existing.tags.values_list("id", flat=True)), prior_tag_ids)
        self.assertEqual(list(existing.tags.values_list("slug", flat=True)), ["invalid"])


class TagManagementTests(TestCase):
    def test_save_creates_tag(self):
        client = staff_client()
        response = client.post(
            reverse("category_type_save"),
            data={"label": "Veterinary", "slug": "vet", "sort_order": "90"},
        )
        self.assertEqual(response.status_code, 302)
        tag = Tag.objects.get(slug="vet")
        self.assertEqual(tag.label, "Veterinary")
        self.assertEqual(tag.sort_order, 90)
        self.assertFalse(tag.is_system)
        self.assertFalse(CategoryRule.objects.filter(category="vet").exists())

    def test_create_seeds_one_rule_per_comma_separated_phrase(self):
        client = staff_client()
        response = client.post(
            reverse("category_type_save"),
            data={
                "label": "Veterinary",
                "sort_order": "90",
                "match_phrase": "Veterinary, Vet Clinic, Animal Hospital, veterinary",
                "priority": "10",
            },
        )
        self.assertEqual(response.status_code, 302)
        tag = Tag.objects.get(slug="veterinary")
        self.assertEqual(tag.label, "Veterinary")
        phrases = list(
            CategoryRule.objects.filter(category="veterinary")
            .order_by("id")
            .values_list("match_phrase", "priority")
        )
        self.assertEqual(
            phrases,
            [
                ("Veterinary", 100),
                ("Vet Clinic", 100),
                ("Animal Hospital", 100),
            ],
        )

        html = client.get(reverse("category_rules")).content.decode()
        self.assertIn("Veterinary, Vet Clinic, Animal Hospital", html)
        self.assertNotIn("Add rule", html)
        self.assertNotIn("Import rules", html)

    def test_update_reconciles_match_phrases_without_resetting_priority(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "Veterinary, Vet Clinic",
            },
        )
        kept = CategoryRule.objects.get(category="vet", match_phrase="Veterinary")
        CategoryRule.objects.filter(pk=kept.pk).update(priority=10)
        kept.refresh_from_db()
        dropped_id = CategoryRule.objects.get(
            category="vet", match_phrase="Vet Clinic"
        ).pk

        tag = Tag.objects.get(slug="vet")
        response = client.post(
            reverse("category_type_save"),
            data={
                "id": str(tag.pk),
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "Veterinary, Animal Hospital",
            },
        )
        self.assertEqual(response.status_code, 302)
        kept.refresh_from_db()
        self.assertEqual(kept.priority, 10)
        self.assertEqual(kept.match_phrase, "Veterinary")
        self.assertFalse(CategoryRule.objects.filter(pk=dropped_id).exists())
        added = CategoryRule.objects.get(category="vet", match_phrase="Animal Hospital")
        self.assertEqual(added.priority, 100)
        self.assertEqual(CategoryRule.objects.filter(category="vet").count(), 2)

    def test_update_empty_match_phrases_deletes_all_rules(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "Veterinary, Vet Clinic",
            },
        )
        tag = Tag.objects.get(slug="vet")
        self.assertEqual(CategoryRule.objects.filter(category="vet").count(), 2)

        response = client.post(
            reverse("category_type_save"),
            data={
                "id": str(tag.pk),
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Tag.objects.filter(pk=tag.pk).exists())
        self.assertFalse(CategoryRule.objects.filter(category="vet").exists())

    def test_manage_list_shows_lead_count_per_tag(self):
        dental = Tag.objects.get(slug="dental")
        gp = Tag.objects.get(slug="gp")
        Lead.objects.create(name="Dental Count A", address="1 Count St").tags.add(dental)
        Lead.objects.create(name="Dental Count B", address="2 Count St").tags.add(dental)
        Lead.objects.create(name="GP Count A", address="3 Count St").tags.add(gp)

        html = staff_client().get(reverse("category_rules")).content.decode()
        self.assertRegex(html, r'data-tag-slug="dental"[^>]*>2 leads<')
        self.assertRegex(html, r'data-tag-slug="gp"[^>]*>1 lead<')
        self.assertRegex(html, r'data-tag-slug="aesthetic"[^>]*>0 leads<')

    def test_manage_page_has_single_save_changes_control(self):
        html = staff_client().get(reverse("category_rules")).content.decode()
        self.assertIn("Save changes", html)
        self.assertIn('id="tags-table-save"', html)
        self.assertIn('data-tag-form="save-all"', html)
        self.assertIn('data-tag-add="1"', html)
        self.assertIn("Add tag", html)
        self.assertIn("Delete", html)
        self.assertNotIn(">Save</button>", html)
        self.assertEqual(html.count("Save changes"), 1)

    def test_bulk_save_updates_multiple_rows(self):
        client = staff_client()
        dental = Tag.objects.get(slug="dental")
        gp = Tag.objects.get(slug="gp")
        response = client.post(
            reverse("category_type_bulk_save"),
            data=_table_save_data(
                {
                    dental.pk: {
                        "label": "Dental clinics",
                        "sort_order": "15",
                        "match_phrase": "Dental, Pergigian",
                    },
                    gp.pk: {
                        "label": "General practice",
                        "match_phrase": "GP clinic",
                    },
                }
            ),
        )
        self.assertEqual(response.status_code, 302)
        dental.refresh_from_db()
        gp.refresh_from_db()
        self.assertEqual(dental.label, "Dental clinics")
        self.assertEqual(dental.sort_order, 15)
        self.assertEqual(gp.label, "General practice")
        self.assertEqual(
            list(
                CategoryRule.objects.filter(category="dental")
                .order_by("id")
                .values_list("match_phrase", flat=True)
            ),
            ["Dental", "Pergigian"],
        )
        self.assertEqual(
            CategoryRule.objects.get(category="gp").match_phrase,
            "GP clinic",
        )

    def test_bulk_save_swaps_slugs_and_rules(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={"label": "Alpha", "slug": "alpha", "sort_order": "200"},
        )
        client.post(
            reverse("category_type_save"),
            data={"label": "Beta", "slug": "beta", "sort_order": "210"},
        )
        alpha = Tag.objects.get(slug="alpha")
        beta = Tag.objects.get(slug="beta")
        CategoryRule.objects.create(
            match_phrase="alpha-rule",
            category="alpha",
            priority=10,
        )
        CategoryRule.objects.create(
            match_phrase="beta-rule",
            category="beta",
            priority=10,
        )
        lead = Lead.objects.create(name="Alpha Swap Clinic", address="1 Swap St")
        lead.tags.add(alpha)

        response = client.post(
            reverse("category_type_bulk_save"),
            data=_table_save_data(
                {
                    alpha.pk: {"slug": "beta", "match_phrase": "alpha-rule"},
                    beta.pk: {"slug": "alpha", "match_phrase": "beta-rule"},
                }
            ),
        )
        self.assertEqual(response.status_code, 302)
        alpha.refresh_from_db()
        beta.refresh_from_db()
        self.assertEqual(alpha.slug, "beta")
        self.assertEqual(beta.slug, "alpha")
        self.assertEqual(
            CategoryRule.objects.get(match_phrase="alpha-rule").category,
            "beta",
        )
        self.assertEqual(
            CategoryRule.objects.get(match_phrase="beta-rule").category,
            "alpha",
        )
        lead.refresh_from_db()
        self.assertEqual(list(lead.tags.values_list("slug", flat=True)), ["beta"])

    def test_bulk_save_rejects_duplicate_slugs(self):
        client = staff_client()
        dental = Tag.objects.get(slug="dental")
        gp = Tag.objects.get(slug="gp")
        response = client.post(
            reverse("category_type_bulk_save"),
            data=_table_save_data(
                {
                    dental.pk: {"slug": "same-slug"},
                    gp.pk: {"slug": "same-slug"},
                }
            ),
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.content.decode(),
            "A tag with this slug already exists.",
        )
        dental.refresh_from_db()
        gp.refresh_from_db()
        self.assertEqual(dental.slug, "dental")
        self.assertEqual(gp.slug, "gp")

    def test_bulk_save_via_fragment_header(self):
        client = staff_client()
        dental = Tag.objects.get(slug="dental")
        response = client.post(
            reverse("category_type_bulk_save"),
            data=_table_save_data({dental.pk: {"label": "Dental fragment"}}),
            HTTP_X_TAGS_FRAGMENT="1",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("Dental fragment", html)
        self.assertIn("Save changes", html)
        self.assertNotIn(">Save</button>", html)
        dental.refresh_from_db()
        self.assertEqual(dental.label, "Dental fragment")

    def test_save_updates_tag_slug_and_derived_category_refs(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={"label": "Pilates", "slug": "pilates", "sort_order": "80"},
        )
        tag = Tag.objects.get(slug="pilates")
        original_pk = tag.pk
        lead = Lead.objects.create(
            name="Pilates Rename Clinic",
            address="1 Rename St",
        )
        lead.tags.add(tag)
        CategoryRule.objects.create(
            match_phrase="pilates",
            category="pilates",
            priority=10,
        )

        response = client.post(
            reverse("category_type_save"),
            data={
                "id": str(tag.pk),
                "label": "Pilates Studio",
                "slug": "pilates_studio",
                "sort_order": "75",
            },
        )
        self.assertEqual(response.status_code, 302)
        tag.refresh_from_db()
        self.assertEqual(tag.pk, original_pk)
        self.assertEqual(tag.slug, "pilates_studio")
        self.assertEqual(tag.label, "Pilates Studio")
        self.assertEqual(tag.sort_order, 75)
        self.assertFalse(Tag.objects.filter(slug="pilates").exists())
        lead.refresh_from_db()
        self.assertEqual(list(lead.tags.values_list("slug", flat=True)), ["pilates_studio"])
        self.assertEqual(
            CategoryRule.objects.get(match_phrase="pilates").category,
            "pilates_studio",
        )

    def test_delete_unused_tag(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={"label": "Pilates", "slug": "pilates", "sort_order": "80"},
        )
        tag = Tag.objects.get(slug="pilates")

        response = client.post(reverse("category_type_delete", kwargs={"pk": tag.pk}))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Tag.objects.filter(pk=tag.pk).exists())
        self.assertFalse(CategoryRule.objects.filter(category="pilates").exists())

    def test_delete_blocked_when_lead_or_rule_still_uses_tag(self):
        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "Veterinary",
            },
        )
        tag = Tag.objects.get(slug="vet")

        lead = Lead.objects.create(
            name="In-use Tag Clinic",
            address="10 Sync St",
        )
        lead.tags.add(tag)
        blocked_lead = client.post(reverse("category_type_delete", kwargs={"pk": tag.pk}))
        self.assertEqual(blocked_lead.status_code, 400)
        self.assertEqual(
            blocked_lead.content.decode(),
            "Leads still use this tag. Reassign them before deleting.",
        )
        self.assertTrue(Tag.objects.filter(pk=tag.pk).exists())
        self.assertTrue(CategoryRule.objects.filter(category="vet").exists())

        lead.tags.clear()
        lead.tags.add(tag)
        blocked_tag = client.post(reverse("category_type_delete", kwargs={"pk": tag.pk}))
        self.assertEqual(blocked_tag.status_code, 400)
        self.assertEqual(
            blocked_tag.content.decode(),
            "Leads still use this tag. Reassign them before deleting.",
        )
        self.assertTrue(Tag.objects.filter(pk=tag.pk).exists())

    @override_settings(SERPER_API_KEY="test-key", HUNT_MAX_LIMIT=100)
    @patch("leads.services.requests.post")
    def test_import_assigns_tag_added_from_manage_page(self, mock_post):
        from leads.services import fetch_leads_from_serper

        client = staff_client()
        client.post(
            reverse("category_type_save"),
            data={
                "label": "Veterinary",
                "slug": "vet",
                "sort_order": "90",
                "match_phrase": "veterinary",
            },
        )
        self.assertTrue(Tag.objects.filter(slug="vet").exists())
        self.assertEqual(
            CategoryRule.objects.get(category="vet").match_phrase,
            "veterinary",
        )

        mock_post.return_value = _serper_places_response(
            [
                {
                    "title": "Sunrise Veterinary Hospital",
                    "address": "20 Sync Import St",
                    "phoneNumber": "+60121111009",
                },
            ]
        )
        result = fetch_leads_from_serper(
            "Johor Bahru",
            "",
            num=20,
            shop_keyword="clinic",
            country="Malaysia",
        )
        self.assertEqual(result.created, 1)
        lead = Lead.objects.get(name="Sunrise Veterinary Hospital")
        self.assertEqual(lead.primary_tag_slug, "vet")
        self.assertEqual(list(lead.tags.values_list("slug", flat=True)), ["vet"])


class TagWritePathTests(TestCase):
    def setUp(self):
        self.group = get_or_create_uncategorized_group()
        self.client = staff_client()
        self.lead = Lead.objects.create(
            name="Tag Write Clinic",
            address="10 Tag Write St",
            group=self.group,
        )
        self.lead.tags.set(Tag.objects.filter(slug="unknown"))

    def test_derived_category_prefers_non_system_then_sort_order(self):
        gp = Tag.objects.get(slug="gp")
        dental = Tag.objects.get(slug="dental")
        unknown = Tag.objects.get(slug="unknown")
        invalid = Tag.objects.get(slug="invalid")
        ordered = Tag.from_slugs(["dental", "gp"])
        self.assertEqual([t.slug for t in ordered], ["dental", "gp"])
        self.assertEqual(Tag.derived_category_slug(ordered), "dental")
        self.assertEqual(
            Tag.derived_category_slug(Tag.from_slugs(["unknown", "dental"])),
            "dental",
        )
        self.assertEqual(Tag.derived_category_slug([unknown]), "unknown")
        self.assertEqual(Tag.derived_category_slug([invalid]), "invalid")
        self.assertEqual(Tag.derived_category_slug([]), "unknown")

    def test_clinic_update_sets_tags_and_derives_category(self):
        response = self.client.patch(
            reverse("clinic_update", kwargs={"pk": self.lead.pk}),
            data=json.dumps(
                {
                    "name": self.lead.name,
                    "address": self.lead.address,
                    "tags": ["dental", "gp"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertCountEqual(payload["tags"], ["dental", "gp"])
        self.assertEqual(payload["category"], "dental")
        self.lead.refresh_from_db()
        self.assertCountEqual(
            list(self.lead.tags.values_list("slug", flat=True)),
            ["dental", "gp"],
        )
        self.assertEqual(self.lead.primary_tag_slug, "dental")

    def test_clinic_update_empty_tags_clears_m2m_and_uses_unknown(self):
        self.lead.tags.set(Tag.objects.filter(slug__in=["dental", "gp"]))
        response = self.client.patch(
            reverse("clinic_update", kwargs={"pk": self.lead.pk}),
            data=json.dumps(
                {
                    "name": self.lead.name,
                    "address": self.lead.address,
                    "tags": [],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.lead.refresh_from_db()
        self.assertEqual(list(self.lead.tags.values_list("slug", flat=True)), [])
        self.assertEqual(self.lead.primary_tag_slug, "unknown")

    def test_clinic_update_rejects_unknown_tag_slug(self):
        response = self.client.patch(
            reverse("clinic_update", kwargs={"pk": self.lead.pk}),
            data=json.dumps(
                {
                    "name": self.lead.name,
                    "address": self.lead.address,
                    "tags": ["not-a-real-tag"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Invalid tag.")

    def test_manual_create_assigns_selected_tags(self):
        response = self.client.post(
            reverse("lead_manual_create"),
            data=json.dumps(
                {
                    "name": "Manual Tag Lead",
                    "address": "11 Tag Write St",
                    "tags": ["aesthetic", "dental"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        lead = Lead.objects.get(pk=response.json()["id"])
        self.assertCountEqual(
            list(lead.tags.values_list("slug", flat=True)),
            ["aesthetic", "dental"],
        )
        self.assertEqual(lead.primary_tag_slug, "dental")

    def test_bulk_manual_replaces_tags_on_owned_leads(self):
        other = Lead.objects.create(
            name="Bulk Other Clinic",
            address="12 Tag Write St",
            group=self.group,
        )
        response = self.client.post(
            reverse("leads_bulk_manual"),
            data=json.dumps(
                {
                    "ids": [self.lead.pk, other.pk],
                    "tags": ["fitness", "cafe"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["updated"], 2)
        self.assertCountEqual(payload["tags"], ["fitness", "cafe"])
        self.assertEqual(payload["category"], "fitness")
        for lead in (self.lead, other):
            lead.refresh_from_db()
            self.assertCountEqual(
                list(lead.tags.values_list("slug", flat=True)),
                ["fitness", "cafe"],
            )
            self.assertEqual(lead.primary_tag_slug, "fitness")
            self.assertTrue(lead.is_processed)

    def test_bulk_manual_requires_at_least_one_tag(self):
        response = self.client.post(
            reverse("leads_bulk_manual"),
            data=json.dumps({"ids": [self.lead.pk], "tags": []}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Select at least one tag.")

    def test_api_list_and_detail_include_tag_slugs(self):
        self.lead.tags.set(Tag.objects.filter(slug__in=["dental", "gp"]))
        detail = self.client.get(f"/api/clinics/{self.lead.pk}")
        self.assertEqual(detail.status_code, 200)
        self.assertCountEqual(detail.json()["tags"], ["dental", "gp"])

        listing = self.client.get("/api/clinics/")
        self.assertEqual(listing.status_code, 200)
        rows = listing.json()
        match = next(row for row in rows if row["id"] == self.lead.pk)
        self.assertCountEqual(match["tags"], ["dental", "gp"])

    def test_dashboard_write_dialogs_use_tag_pickers(self):
        html = self.client.get(reverse("dashboard")).content.decode()
        self.assertIn('id="clinic-edit-tags"', html)
        self.assertIn('id="lead-create-tags"', html)
        self.assertIn('id="bulk-manual-tags"', html)
        self.assertNotIn('id="clinic-edit-type"', html)
        self.assertNotIn('id="lead-create-type"', html)
        self.assertNotIn('id="bulk-manual-category"', html)
        self.assertIn(">Set tags<", html)
        self.assertNotIn(">Set category<", html)

    def test_edit_dialog_js_prefills_from_data_tags(self):
        js = (
            Path(__file__).resolve().parent
            / "static"
            / "leads"
            / "js"
            / "lead-dialogs.js"
        ).read_text(encoding="utf-8")
        self.assertIn("Array.isArray(data.tags)", js)
        self.assertIn("clinic-edit-tags", js)
        self.assertIn("setTagPickerSlugs", js)
        self.assertNotIn("clinic-edit-type", js)


class TagBackupAndScriptGroupTests(TestCase):
    def test_backup_round_trip_preserves_tags_on_fresh_dataset(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from leads.backup import build_backup_workbook, restore_from_workbook

        vet = Tag.objects.create(slug="veterinary", label="Veterinary", sort_order=90)
        gp = Tag.objects.get(slug="gp")
        lead = Lead.objects.create(
            name="Sunrise Veterinary Hospital",
            address="20 Backup St",
        )
        lead.tags.set([vet, gp])

        wb = build_backup_workbook()
        buf = BytesIO()
        wb.save(buf)
        payload = buf.getvalue()

        loaded = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        self.assertIn("Tags", loaded.sheetnames)
        headers = [cell.value for cell in next(loaded["Leads"].iter_rows(min_row=1, max_row=1))]
        self.assertIn("tags", headers)
        tag_idx = headers.index("tags")
        lead_row = next(
            row
            for row in loaded["Leads"].iter_rows(min_row=2, values_only=True)
            if row[1] == "Sunrise Veterinary Hospital"
        )
        self.assertCountEqual(
            [part.strip() for part in str(lead_row[tag_idx]).split(",") if part.strip()],
            ["veterinary", "gp"],
        )

        lead.delete()
        vet.delete()
        self.assertFalse(Tag.objects.filter(slug="veterinary").exists())
        self.assertFalse(
            Lead.objects.filter(
                name="Sunrise Veterinary Hospital", address="20 Backup St"
            ).exists()
        )

        summary = restore_from_workbook(BytesIO(payload))
        self.assertEqual(summary["leads_created"], 1)
        restored = Lead.objects.get(
            name="Sunrise Veterinary Hospital", address="20 Backup St"
        )
        self.assertCountEqual(
            list(restored.tags.values_list("slug", flat=True)),
            ["veterinary", "gp"],
        )
        self.assertEqual(Tag.objects.get(slug="veterinary").label, "Veterinary")
        self.assertEqual(
            restored.primary_tag_slug,
            Tag.derived_category_slug(list(restored.tags.all())),
        )

    def test_legacy_backup_without_tags_column_still_assigns_a_tag(self):
        from io import BytesIO

        from openpyxl import Workbook

        from leads.backup import LEAD_HEADERS, restore_from_workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        old_headers = [header for header in LEAD_HEADERS if header != "tags"]
        ws.append(old_headers)
        values = {header: "" for header in old_headers}
        values.update(
            {
                "id": 99,
                "name": "Legacy Dental Clinic",
                "address": "1 Old Backup St",
                "category": "dental",
            }
        )
        ws.append([values[header] for header in old_headers])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        summary = restore_from_workbook(buf)
        self.assertEqual(summary["leads_created"], 1)
        lead = Lead.objects.get(name="Legacy Dental Clinic", address="1 Old Backup St")
        self.assertEqual(lead.primary_tag_slug, "dental")
        self.assertEqual(list(lead.tags.values_list("slug", flat=True)), ["dental"])

    def test_restore_does_not_change_tags_on_existing_leads(self):
        from io import BytesIO

        from leads.backup import build_backup_workbook, restore_from_workbook

        gp = Tag.objects.get(slug="gp")
        dental = Tag.objects.get(slug="dental")
        lead = Lead.objects.create(
            name="Stay Put Clinic",
            address="9 Skip St",
        )
        lead.tags.set([gp])
        wb = build_backup_workbook()
        buf = BytesIO()
        wb.save(buf)

        lead.tags.set([dental])

        summary = restore_from_workbook(BytesIO(buf.getvalue()))
        self.assertEqual(summary["leads_skipped"], 1)
        lead.refresh_from_db()
        self.assertEqual(list(lead.tags.values_list("slug", flat=True)), ["dental"])

    def test_script_group_uses_tag_label_not_hardcoded_category_display(self):
        from leads.pipeline import get_or_create_uncategorized_group
        from leads.whatsapp_service import (
            SCRIPT_TEMPLATE_FALLBACK_GROUP,
            script_group_name_for_lead,
        )

        tag = Tag.objects.create(
            slug="psychologist",
            label="Psychologist",
            sort_order=95,
        )
        lead = Lead.objects.create(
            name="Mind Clinic",
            address="1 Psy St",
            group=get_or_create_uncategorized_group(),
        )
        lead.tags.set([tag])

        self.assertFalse(Tag.objects.filter(slug="psychologist", is_system=True).exists())
        from leads.category_types import DEFAULT_CATEGORY_TYPES

        self.assertNotIn("psychologist", [row[0] for row in DEFAULT_CATEGORY_TYPES])
        self.assertEqual(script_group_name_for_lead(lead), "Psychologist")

        gp_lead = Lead.objects.create(
            name="Town GP",
            address="2 Psy St",
            group=get_or_create_uncategorized_group(),
        )
        gp_lead.tags.set(Tag.objects.filter(slug="gp"))
        self.assertEqual(script_group_name_for_lead(gp_lead), "GP")

        unknown = Lead.objects.create(
            name="Unknown Shop",
            address="3 Psy St",
            group=get_or_create_uncategorized_group(),
        )
        unknown.tags.set(Tag.objects.filter(slug="unknown"))
        self.assertEqual(script_group_name_for_lead(unknown), SCRIPT_TEMPLATE_FALLBACK_GROUP)
