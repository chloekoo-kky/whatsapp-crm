"""Import leads from the dashboard “Export to Excel” workbook."""

from __future__ import annotations

from datetime import datetime

from django.db import transaction
from django.db.models import Max

from leads.backup import (
    _make_aware,
    _p_bool,
    _p_dt,
    _p_str,
    _sheet,
    _tags_and_category_for_row,
)
from leads.display import normalize_manual_phone
from leads.models import Lead
from leads.pipeline import QUEUE_GROUP_NAME
from leads.services import sync_chain_flags_for_name

MAX_PHONES_PER_LEAD = 8

# Display headers from clinics_export_xlsx, plus a few backup-style aliases.
_HEADER_ALIASES = {
    "name": "name",
    "phone": "phone",
    "address": "address",
    "website": "website",
    "keyword": "shop_keyword",
    "shop_keyword": "shop_keyword",
    "tags": "tags",
    "search city": "search_city",
    "search_city": "search_city",
    "search state": "search_state",
    "search_state": "search_state",
    "search country": "search_country",
    "search_country": "search_country",
    "maps query": "search_query",
    "search_query": "search_query",
    "chain": "is_chain",
    "is_chain": "is_chain",
    "very important": "is_very_important",
    "is_very_important": "is_very_important",
    "processed": "is_processed",
    "is_processed": "is_processed",
    "source url": "source_url",
    "source_url": "source_url",
    "created": "created_at",
    "created_at": "created_at",
}

_BACKUP_SHEETS = frozenset(
    {"ChatMessages", "WhatsAppConfig", "ScriptTemplates", "ConversationLogs"}
)


def _header_key(raw) -> str:
    return " ".join(_p_str(raw).strip().lower().split())


def _remap_row(row: dict) -> dict:
    mapped: dict = {}
    for key, val in row.items():
        alias = _HEADER_ALIASES.get(_header_key(key))
        if alias:
            mapped[alias] = val
    return mapped


def _parse_phones(raw) -> list[str]:
    text = _p_str(raw).strip()
    if not text:
        return []
    out: list[str] = []
    for chunk in text.replace(",", ";").split(";"):
        n = normalize_manual_phone(chunk.strip())
        if n and n not in out:
            out.append(n[:64])
        if len(out) >= MAX_PHONES_PER_LEAD:
            break
    return out


def _parse_created(value):
    dt = _p_dt(value)
    if dt:
        return dt
    text = _p_str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return _make_aware(datetime.strptime(text, fmt))
        except ValueError:
            continue
    return None


@transaction.atomic
def import_exported_leads_workbook(file_obj, *, group, assigned_to=None) -> dict:
    """Create leads from an Export to Excel workbook. Skip existing name+address."""
    from openpyxl import load_workbook

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    if _BACKUP_SHEETS.intersection(wb.sheetnames):
        raise ValueError(
            "This looks like a full backup file. Use Restore instead of Import from Excel."
        )

    ws = _sheet(wb, "Leads") or wb.active
    if ws is None:
        raise ValueError("No worksheet found in the Excel file.")

    rows = ws.iter_rows(values_only=True)
    raw_headers = next(rows, None)
    if not raw_headers:
        raise ValueError(
            "Not an Export to Excel file. Expected a worksheet with a Name column."
        )
    headers = [_p_str(h).strip() for h in raw_headers]
    if "name" not in {_HEADER_ALIASES.get(_header_key(h)) for h in headers}:
        raise ValueError(
            "Not an Export to Excel file. Expected a worksheet with a Name column."
        )

    summary = {"leads_created": 0, "leads_skipped": 0, "rows_ignored": 0}
    next_order = (Lead.objects.filter(group=group).aggregate(m=Max("display_order"))["m"] or 0) + 1
    created_names: set[str] = set()
    queue_group = group is not None and group.name == QUEUE_GROUP_NAME

    for values in rows:
        if values is None or all(cell is None or cell == "" for cell in values):
            continue
        row = _remap_row(dict(zip(headers, values)))
        name = _p_str(row.get("name")).strip()
        if not name:
            summary["rows_ignored"] += 1
            continue
        address = _p_str(row.get("address"))
        phones = _parse_phones(row.get("phone"))
        status = Lead.WhatsappStatus.IDLE
        if queue_group and phones:
            status = Lead.WhatsappStatus.PENDING
        defaults = {
            "phone_number": phones[0] if phones else "",
            "phone_numbers": phones,
            "website": _p_str(row.get("website")).strip()[:500],
            "shop_keyword": _p_str(row.get("shop_keyword")).strip()[:160],
            "source_url": _p_str(row.get("source_url")).strip(),
            "is_processed": _p_bool(row.get("is_processed")),
            "is_chain": _p_bool(row.get("is_chain")),
            "is_very_important": _p_bool(row.get("is_very_important")),
            "search_city": _p_str(row.get("search_city")).strip() or None,
            "search_state": _p_str(row.get("search_state")).strip() or None,
            "search_query": _p_str(row.get("search_query")).strip() or None,
            "search_country": _p_str(row.get("search_country")).strip() or None,
            "group": group,
            "whatsapp_status": status,
            "display_order": next_order,
            "assigned_to": assigned_to,
        }
        lead, created = Lead.objects.get_or_create(
            name=name[:255],
            address=address,
            defaults=defaults,
        )
        if not created:
            summary["leads_skipped"] += 1
            continue
        summary["leads_created"] += 1
        next_order += 1
        created_names.add(lead.name)
        created_at = _parse_created(row.get("created_at"))
        if created_at:
            Lead.objects.filter(pk=lead.pk).update(created_at=created_at)
        tags, _derived = _tags_and_category_for_row(row)
        if tags:
            lead.tags.set(tags)

    for name in created_names:
        sync_chain_flags_for_name(name)

    return summary
